"""Handelsübersicht als Excel (Übersicht · Nach Verkaufsort · Nach Kaufort · Alle Preise · Info).

Logik aus update_starhead.py hierher verschoben, damit Web-Build (GitHub Actions) und
lokale Pipeline dieselbe Datei erzeugen. Braucht openpyxl.
"""
from handel_core import by_buy_location, by_sell_location


def build_xlsx(rows, summary, stamp, patch, out_path):  # noqa: C901 - lineare Blatt-Formatierung
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    head_fill = PatternFill("solid", fgColor="0B3D5C")
    head_font = Font(color="FFFFFF", bold=True, size=10)
    ill = PatternFill("solid", fgColor="FBE4E4")
    pos = PatternFill("solid", fgColor="E7F4EA")
    num = "#,##0"
    green = Font(bold=True, color="1B6B2E")
    red = Font(color="B00020", bold=True)

    def hdr(ws, n):
        for c in range(1, n + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(n)}{ws.max_row}"
        ws.row_dimensions[1].height = 26

    def widths(ws, ws_widths):
        for i, w in enumerate(ws_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    def numfmt(ws, cols):
        for r in range(2, ws.max_row + 1):
            for c in cols:
                cell = ws.cell(row=r, column=c)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = num

    wb = Workbook()
    ws = wb.active
    ws.title = "Übersicht"
    cols = ["Ware", "Typ", "Illegal", "Basispreis", "Bester Kauf (aUEC/SCU)", "Kaufort", "System (Kauf)",
            "Kauf-Bestand SCU", "Bester Verkauf (aUEC/SCU)", "Verkaufsort", "System (Verkauf)", "Marge/SCU",
            "#Kauforte", "#Verkaufsorte"]
    ws.append(cols)

    def mkey(s):
        m = s.get("marge")
        return (0, -m) if isinstance(m, (int, float)) else (1, 0)

    for s in sorted(summary, key=mkey):
        ws.append([s["ware"], s["typ"], "Ja" if s["illegal"] else "", s.get("basePrice", ""), s.get("kaufPreis", ""),
                   s.get("kaufOrt", ""), s.get("kaufSys", ""), s.get("kaufInv", ""), s.get("verkaufPreis", ""),
                   s.get("verkaufOrt", ""), s.get("verkaufSys", ""), s.get("marge", ""), s.get("nKauf", ""),
                   s.get("nVerkauf", "")])
    hdr(ws, len(cols))
    widths(ws, [22, 16, 7, 10, 20, 30, 15, 14, 22, 34, 16, 12, 10, 13])
    numfmt(ws, [4, 5, 8, 9, 12])
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=3).value == "Ja":
            for c in range(1, len(cols) + 1):
                ws.cell(row=r, column=c).fill = ill
        mc = ws.cell(row=r, column=12)
        if isinstance(mc.value, (int, float)) and mc.value > 0:
            mc.font = green

    # ---------- gruppierte Blätter ----------
    ga = PatternFill("solid", fgColor="FFFFFF")
    gb = PatternFill("solid", fgColor="EEF3F6")

    def grouped(title, cols_, recs, key, rowfn, ws_widths, numcols, margecol, illcol):
        wsg = wb.create_sheet(title)
        wsg.append(cols_)
        prev, gi = None, -1
        for x in recs:
            if x[key] != prev:
                gi += 1
                prev = x[key]
            wsg.append(rowfn(x))
            fill = gb if gi % 2 else ga
            for c in range(1, len(cols_) + 1):
                wsg.cell(row=wsg.max_row, column=c).fill = fill
        hdr(wsg, len(cols_))
        widths(wsg, ws_widths)
        numfmt(wsg, numcols)
        for r in range(2, wsg.max_row + 1):
            mc = wsg.cell(row=r, column=margecol)
            if isinstance(mc.value, (int, float)) and mc.value > 0:
                mc.font = green
            if wsg.cell(row=r, column=illcol).value == "Ja":
                wsg.cell(row=r, column=illcol).font = red

    grouped("Nach Verkaufsort",
            ["Verkaufsort", "System (Verkauf)", "Ware", "Typ", "Illegal", "Verkauf (aUEC/SCU)", "Bester Kauf",
             "Kaufort", "System (Kauf)", "Marge/SCU", "Stand"],
            by_sell_location(rows, summary), "vort",
            lambda x: [x["vort"], x["sys"], x["ware"], x["typ"], "Ja" if x["illegal"] else "", x["vpreis"],
                       x["kpreis"], x["kort"], x["ksys"], x["marge"], x["stand"]],
            [32, 15, 22, 15, 7, 17, 14, 30, 13, 12, 11], [6, 7, 10], 10, 5)
    grouped("Nach Kaufort",
            ["Kaufort", "System (Kauf)", "Ware", "Typ", "Illegal", "Kauf (aUEC/SCU)", "Bester Verkauf",
             "Verkaufsort", "System (Verkauf)", "Marge/SCU", "Kauf-Bestand SCU", "Stand"],
            by_buy_location(rows, summary), "kort",
            lambda x: [x["kort"], x["sys"], x["ware"], x["typ"], "Ja" if x["illegal"] else "", x["kpreis"],
                       x["vpreis"], x["vort"], x["vsys"], x["marge"], x.get("inv", ""), x["stand"]],
            [30, 13, 22, 15, 7, 15, 15, 32, 14, 12, 15, 11], [6, 7, 10, 11], 10, 5)

    # ---------- Alle Preise ----------
    ws2 = wb.create_sheet("Alle Preise")
    cols2 = ["Ware", "Typ", "Illegal", "Art", "Preis (aUEC/SCU)", "Shop", "Ort", "System", "Max-Inventar (SCU)",
             "Container (SCU)", "Stand"]
    ws2.append(cols2)
    for x in sorted(rows, key=lambda r: (r["ware"].lower(), 0 if r["art"] == "Kauf" else 1, r["preis"])):
        ws2.append([x["ware"], x["typ"], "Ja" if x["illegal"] else "", x["art"], x["preis"], x["shop"], x["ort"],
                    x["system"], x.get("invScu", ""), x.get("container", ""), x.get("stand", "")])
    hdr(ws2, len(cols2))
    widths(ws2, [22, 16, 7, 9, 16, 32, 18, 16, 17, 14, 11])
    numfmt(ws2, [5])
    for r in range(2, ws2.max_row + 1):
        if ws2.cell(row=r, column=4).value == "Verkauf":
            ws2.cell(row=r, column=4).fill = pos
        if ws2.cell(row=r, column=3).value == "Ja":
            ws2.cell(row=r, column=3).font = red

    # ---------- Info ----------
    ws3 = wb.create_sheet("Info")
    for row in [["SC Waren-Handelsübersicht", ""], ["Quelle", "star-head.de via API (commodity/shop/shopitem-price)"],
                ["Abgerufen am", stamp], ["Spielstand", "Star Citizen " + patch], ["Anzahl Waren", len(summary)],
                ["Anzahl Preiszeilen", len(rows)], ["", ""], ["Hinweise", ""], ["Preise", "aUEC pro SCU."],
                ["Marge/SCU", "Bester Verkauf − bester Kauf (grün = positiv). Ohne Reise-/Zeitkosten."],
                ["Kauf-Bestand", "Bester Kauf = günstigster Ort mit Bestand>0 (sofern vorhanden)."],
                ["Aktualität", "'Stand' = letztes Community-Update; teils Tage/Wochen alt."],
                ["Illegal", "Rot: nur an versteckten Konsolen/Outlaw-Stationen; Scan-Risiko."],
                ["Tab 'Nach Verkaufsort'", "Je Verkaufsort gruppiert (beste Marge oben), darin Waren nach Marge absteigend. "
                 "Zeigt, was man an einem Ort alles profitabel verkaufen kann (Mehr-Waren-Ladung planen)."],
                ["Tab 'Nach Kaufort'", "Je Kaufort gruppiert (beste Marge oben), darin Waren nach Marge absteigend. "
                 "Zeigt, was man an einem Ort mit guter Marge einladen kann (inkl. bestem Verkaufsort + Kauf-Bestand)."],
                ["Online", "https://jaegerfeld.github.io/StarCitizen-Handels-Planer/ (täglich neu gebaut)"],
                ["Auto-Update", "Neu erzeugt via update_starhead.py / handels_planer_bauen.py"]]:
        ws3.append(row)
    ws3.column_dimensions["A"].width = 20
    ws3.column_dimensions["B"].width = 95
    ws3["A1"].font = Font(bold=True, size=14, color="0B3D5C")
    for r in range(2, ws3.max_row + 1):
        ws3.cell(row=r, column=1).font = Font(bold=True)
        ws3.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(out_path)
    return out_path
